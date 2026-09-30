"""Integration tests for the preliminary and final Excel contract."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from asset_servicing.adapters.export import ExcelExporter, ExcelExportError
from asset_servicing.adapters.persistence import JsonRunRepository, RunEvent
from asset_servicing.domain import (
    CoverageFinding,
    ExtractedVariable,
    FindingReviewStatus,
    ReviewAction,
    ReviewDecision,
    ReviewStatus,
    Run,
    RunState,
    SectionLocation,
    SourceKind,
    ValidationResult,
    ValidationVerdict,
)

pytestmark = pytest.mark.integration


CREATED_AT = datetime(2026, 9, 28, 12, 0, tzinfo=UTC)
REVIEWED_AT = datetime(2026, 9, 28, 15, 0, tzinfo=UTC)


def make_run(
    *,
    confidence: float = 0.97,
    verdict: ValidationVerdict = ValidationVerdict.SUPPORTED,
    reviewed: bool = False,
    review_status: ReviewStatus = ReviewStatus.PENDING,
    source_kind: SourceKind = SourceKind.PROSE,
) -> Run:
    return Run(
        run_id="run-001",
        document_name="regulamento.pdf",
        document_sha256="c" * 64,
        page_count=12,
        created_at=CREATED_AT,
        state=RunState.REVIEWING,
        location=SectionLocation(
            title="Emissão, aplicação e resgate",
            page_start=7,
            page_end=9,
            rationale="Intervalo confirmado.",
            confirmed=True,
        ),
        variables=[
            ExtractedVariable(
                id="variable-001",
                canonical_name="prazo_resgate",
                original_name="prazo_resgate",
                original_value="D+30",
                current_name="prazo_pagamento_resgate" if reviewed else "prazo_resgate",
                current_value="até D+30" if reviewed else "D+30",
                evidence_text="O pagamento do resgate ocorrerá em até trinta dias.",
                source_pages=[8],
                source_kind=source_kind,
                reviewed=reviewed,
                review_status=review_status,
            )
        ],
        validations=[
            ValidationResult(
                variable_id="variable-001",
                confidence=confidence,
                verdict=verdict,
                rationale="A fonte sustenta o resultado.",
                issues=[] if verdict is ValidationVerdict.SUPPORTED else ["Revisar contexto"],
            )
        ],
    )


def persist_run(tmp_path: Path, run: Run) -> tuple[JsonRunRepository, ExcelExporter]:
    repository = JsonRunRepository(tmp_path / "runs")
    repository.save_run(run)
    return repository, ExcelExporter(repository)


def add_events(repository: JsonRunRepository) -> None:
    repository.append_event(
        RunEvent(
            run_id="run-001",
            stage="extraction",
            event="stage_completed",
            status="success",
            started_at=CREATED_AT,
            duration_ms=1250,
            model="extractor-model",
            prompt_version="extractor-v1",
            input_pages=[7, 8, 9],
            usage={"total_tokens": 321, "attempts": 1},
        )
    )


def test_preliminary_export_creates_a_downloadable_workbook(tmp_path: Path) -> None:
    repository, exporter = persist_run(
        tmp_path,
        make_run(confidence=0.72, verdict=ValidationVerdict.PARTIALLY_SUPPORTED),
    )

    path = exporter.export_preliminary(repository.load_run("run-001"))

    assert path.is_file()
    assert path.name == "preliminary.xlsx"
    assert repository.load_run("run-001").preliminary_export_path == str(path)


def test_final_export_is_blocked_while_review_is_pending(tmp_path: Path) -> None:
    repository, exporter = persist_run(
        tmp_path,
        make_run(confidence=0.72, verdict=ValidationVerdict.PARTIALLY_SUPPORTED),
    )

    with pytest.raises(ExcelExportError, match="pending"):
        exporter.export_final(repository.load_run("run-001"))

    assert repository.load_run("run-001").final_export_path is None


def test_final_export_is_released_after_human_review(tmp_path: Path) -> None:
    run = make_run(reviewed=True, review_status=ReviewStatus.CONFIRMED)
    run.reviews = [
        ReviewDecision(
            variable_id="variable-001",
            action=ReviewAction.CONFIRM,
            previous_name="prazo_resgate",
            previous_value="D+30",
            resulting_name="prazo_resgate",
            resulting_value="D+30",
            reviewed_at=REVIEWED_AT,
        )
    ]
    repository, exporter = persist_run(tmp_path, run)

    path = exporter.export_final(repository.load_run("run-001"))

    assert path.is_file()
    assert repository.load_run("run-001").final_export_path == str(path)


def test_final_export_audits_a_dismissed_omission(tmp_path: Path) -> None:
    run = make_run()
    run.coverage_findings = [
        CoverageFinding(
            id="finding-001",
            description="Possível regra do capítulo seguinte.",
            evidence_text="CAPÍTULO 4 - PRESTADORES DE SERVIÇOS",
            source_pages=[9],
            review_status=FindingReviewStatus.DISMISSED,
        )
    ]
    run.reviews = [
        ReviewDecision(
            finding_id="finding-001",
            action=ReviewAction.DISMISS_OMISSION,
            note="Fora da seção-alvo.",
            reviewed_at=REVIEWED_AT,
        )
    ]
    repository, exporter = persist_run(tmp_path, run)

    path = exporter.export_final(repository.load_run("run-001"))
    workbook = load_workbook(path)
    try:
        audit_values = list(workbook["Auditoria"].values)
        assert any("finding-001" in str(value) for row in audit_values for value in row)
        assert any("dismiss_omission" in str(value) for row in audit_values for value in row)
        assert any("Fora da seção-alvo." in str(value) for row in audit_values for value in row)
    finally:
        workbook.close()


def test_main_sheet_has_exact_case_columns_in_order(tmp_path: Path) -> None:
    repository, exporter = persist_run(tmp_path, make_run())
    path = exporter.export_preliminary(repository.load_run("run-001"))

    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook.sheetnames == ["Variáveis", "Auditoria"]
        assert [cell.value for cell in workbook["Variáveis"][1]] == [
            "Variável",
            "Valor da variável",
            "Trecho da variável",
            "Grau de Confiança",
            "Foi revisado?",
        ]
    finally:
        workbook.close()


def test_main_sheet_preserves_numeric_confidence_and_boolean_review(tmp_path: Path) -> None:
    repository, exporter = persist_run(tmp_path, make_run(confidence=0.91))
    path = exporter.export_preliminary(repository.load_run("run-001"))

    workbook = load_workbook(path, data_only=False)
    try:
        row = list(workbook["Variáveis"].iter_rows(min_row=2, values_only=True))[0]
        assert row[3] == 0.91
        assert isinstance(row[3], float)
        assert row[4] is False
        assert isinstance(row[4], bool)
    finally:
        workbook.close()


def test_not_applicable_is_omitted_from_main_but_kept_in_audit(tmp_path: Path) -> None:
    run = make_run(reviewed=True, review_status=ReviewStatus.NOT_APPLICABLE)
    run.reviews = [
        ReviewDecision(
            variable_id="variable-001",
            action=ReviewAction.NOT_APPLICABLE,
            previous_name="prazo_resgate",
            previous_value="D+30",
            resulting_name="prazo_resgate",
            resulting_value="D+30",
            note="Não se aplica.",
            reviewed_at=REVIEWED_AT,
        )
    ]
    repository, exporter = persist_run(tmp_path, run)
    path = exporter.export_preliminary(repository.load_run("run-001"))

    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook["Variáveis"].max_row == 1
        audit_values = list(workbook["Auditoria"].values)
        assert any("not_applicable" in str(value) for row in audit_values for value in row)
    finally:
        workbook.close()


def test_edited_variable_exports_current_name_and_value(tmp_path: Path) -> None:
    run = make_run(reviewed=True, review_status=ReviewStatus.EDITED)
    run.reviews = [
        ReviewDecision(
            variable_id="variable-001",
            action=ReviewAction.EDIT,
            previous_name="prazo_resgate",
            previous_value="D+30",
            resulting_name="prazo_pagamento_resgate",
            resulting_value="até D+30",
            reviewed_at=REVIEWED_AT,
        )
    ]
    repository, exporter = persist_run(tmp_path, run)
    path = exporter.export_preliminary(repository.load_run("run-001"))

    workbook = load_workbook(path, data_only=False)
    try:
        row = list(workbook["Variáveis"].iter_rows(min_row=2, values_only=True))[0]
        assert row[0:2] == ("prazo_pagamento_resgate", "até D+30")
        assert row[2] == "O pagamento do resgate ocorrerá em até trinta dias."
    finally:
        workbook.close()


def test_reopened_workbook_preserves_values_types_and_tabs(tmp_path: Path) -> None:
    repository, exporter = persist_run(tmp_path, make_run())
    path = exporter.export_preliminary(repository.load_run("run-001"))

    reopened = load_workbook(path, data_only=False)
    try:
        assert reopened.sheetnames == ["Variáveis", "Auditoria"]
        assert reopened["Variáveis"]["A2"].value == "prazo_resgate"
        assert reopened["Variáveis"]["D2"].data_type == "n"
        assert reopened["Variáveis"]["E2"].data_type == "b"
    finally:
        reopened.close()


def test_audit_sheet_contains_run_document_hash_pages_and_state(tmp_path: Path) -> None:
    repository, exporter = persist_run(tmp_path, make_run())
    path = exporter.export_preliminary(repository.load_run("run-001"))

    workbook = load_workbook(path, data_only=False)
    try:
        values = [value for row in workbook["Auditoria"].values for value in row]
        assert "run-001" in values
        assert "regulamento.pdf" in values
        assert "c" * 64 in values
        assert "7-9" in values
        assert "reviewing" in values
    finally:
        workbook.close()


def test_audit_sheet_contains_event_model_prompt_duration_and_usage(tmp_path: Path) -> None:
    repository, exporter = persist_run(tmp_path, make_run())
    add_events(repository)
    path = exporter.export_preliminary(repository.load_run("run-001"))

    workbook = load_workbook(path, data_only=False)
    try:
        values = [value for row in workbook["Auditoria"].values for value in row]
        assert "extractor-model" in values
        assert "extractor-v1" in values
        assert 1250 in values
        assert 321 in values
    finally:
        workbook.close()


def test_audit_sheet_contains_review_history(tmp_path: Path) -> None:
    run = make_run(reviewed=True, review_status=ReviewStatus.EDITED)
    run.reviews = [
        ReviewDecision(
            variable_id="variable-001",
            action=ReviewAction.EDIT,
            previous_name="prazo_resgate",
            previous_value="D+30",
            resulting_name="prazo_pagamento_resgate",
            resulting_value="até D+30",
            note="Ajuste humano.",
            reviewed_at=REVIEWED_AT,
        )
    ]
    repository, exporter = persist_run(tmp_path, run)
    path = exporter.export_final(repository.load_run("run-001"))

    workbook = load_workbook(path, data_only=False)
    try:
        values = [value for row in workbook["Auditoria"].values for value in row]
        assert "edit" in values
        assert "Ajuste humano." in values
        assert "até D+30" in values
    finally:
        workbook.close()


def test_export_requires_at_least_one_validated_variable(tmp_path: Path) -> None:
    run = make_run()
    run.variables = []
    run.validations = []
    repository, exporter = persist_run(tmp_path, run)

    with pytest.raises(ExcelExportError, match="validated"):
        exporter.export_preliminary(repository.load_run("run-001"))


def test_preliminary_and_final_exports_use_distinct_paths(tmp_path: Path) -> None:
    run = make_run(reviewed=True, review_status=ReviewStatus.CONFIRMED)
    repository, exporter = persist_run(tmp_path, run)

    preliminary = exporter.export_preliminary(repository.load_run("run-001"))
    final = exporter.export_final(repository.load_run("run-001"))

    assert preliminary != final
    assert preliminary.name == "preliminary.xlsx"
    assert final.name == "final.xlsx"


def test_human_added_variable_is_exported_with_review_marker(tmp_path: Path) -> None:
    run = make_run(
        reviewed=True,
        review_status=ReviewStatus.CONFIRMED,
        source_kind=SourceKind.HUMAN_ADDED,
    )
    run.validations = []
    repository, exporter = persist_run(tmp_path, run)

    path = exporter.export_final(repository.load_run("run-001"))

    workbook = load_workbook(path, data_only=False)
    try:
        row = list(workbook["Variáveis"].iter_rows(min_row=2, values_only=True))[0]
        assert row[0] == "prazo_pagamento_resgate"
        assert row[3] == 1.0
        assert row[4] is True
    finally:
        workbook.close()
