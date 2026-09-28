"""Generate auditable preliminary and final Excel workbooks."""

from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from asset_servicing.adapters.persistence import JsonRunRepository, RunEvent
from asset_servicing.domain import ReviewStatus, Run, SourceKind

MAIN_HEADERS = [
    "Variável",
    "Valor da variável",
    "Trecho da variável",
    "Grau de Confiança",
    "Foi revisado?",
]
AUDIT_HEADERS = ["Categoria", "Campo", "Valor"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


class ExcelExportError(ValueError):
    """Raised when an export is not valid for the current run state."""


class ExcelExporter:
    """Create workbooks from a persisted run and its audit events."""

    def __init__(self, repository: JsonRunRepository) -> None:
        self._repository = repository

    def export_preliminary(self, run: Run) -> Path:
        """Write a preliminary workbook even when review items remain open."""

        return self._export(run, filename="preliminary.xlsx", require_clear=False)

    def export_final(self, run: Run) -> Path:
        """Write the final workbook only when no review item remains pending."""

        return self._export(run, filename="final.xlsx", require_clear=True)

    def _export(self, run: Run, *, filename: str, require_clear: bool) -> Path:
        if require_clear and run.pending_items():
            raise ExcelExportError("final export is blocked while review items are pending")
        rows = self._variable_rows(run)
        has_human_added = any(
            variable.source_kind is SourceKind.HUMAN_ADDED for variable in run.variables
        )
        if not rows and not run.validations and not has_human_added:
            raise ExcelExportError("at least one validated variable is required")

        workbook = Workbook()
        main = workbook.active
        if main is None:
            raise ExcelExportError("workbook has no active worksheet")
        main.title = "Variáveis"
        main.append(MAIN_HEADERS)
        for row in rows:
            main.append(row)
        audit = workbook.create_sheet("Auditoria")
        audit.append(AUDIT_HEADERS)
        for row in self._audit_rows(run):
            audit.append(row)
        self._format_sheet(main)
        self._format_sheet(audit)

        output_dir = self._repository.root / run.run_id / "exports"
        output_dir.mkdir(parents=True, exist_ok=True)
        path = output_dir / filename
        workbook.save(path)
        workbook.close()
        if filename == "preliminary.xlsx":
            run.preliminary_export_path = str(path)
        else:
            run.final_export_path = str(path)
        self._repository.save_run(run)
        return path

    def _variable_rows(self, run: Run) -> list[list[object]]:
        validation_by_id = {item.variable_id: item for item in run.validations}
        rows: list[list[object]] = []
        for variable in run.variables:
            if variable.review_status is ReviewStatus.NOT_APPLICABLE:
                continue
            validation = validation_by_id.get(variable.id)
            if validation is None and variable.source_kind is not SourceKind.HUMAN_ADDED:
                continue
            confidence = 1.0 if validation is None else validation.confidence
            rows.append(
                [
                    variable.current_name,
                    variable.current_value,
                    variable.evidence_text,
                    float(confidence),
                    bool(variable.reviewed),
                ]
            )
        return rows

    def _audit_rows(self, run: Run) -> list[list[object]]:
        rows: list[list[object]] = [
            ["run", "run_id", run.run_id],
            ["run", "document_name", run.document_name],
            ["run", "document_sha256", run.document_sha256],
            ["run", "confirmed_pages", self._confirmed_pages(run)],
            ["run", "state", run.state.value],
            ["run", "created_at", run.created_at.isoformat()],
        ]
        rows.extend(self._event_rows(self._repository.load_events(run.run_id)))
        for decision in run.reviews:
            rows.extend(
                [
                    ["review", "variable_id", decision.variable_id],
                    ["review", "action", decision.action.value],
                    ["review", "previous_name", decision.previous_name],
                    ["review", "previous_value", decision.previous_value],
                    ["review", "resulting_name", decision.resulting_name],
                    ["review", "resulting_value", decision.resulting_value],
                    ["review", "note", decision.note],
                    ["review", "reviewed_at", decision.reviewed_at.isoformat()],
                ]
            )
        return rows

    @staticmethod
    def _event_rows(events: Iterable[RunEvent]) -> list[list[object]]:
        rows: list[list[object]] = []
        for index, event in enumerate(events, start=1):
            rows.extend(
                [
                    [f"event:{index}", "stage", event.stage],
                    [f"event:{index}", "status", event.status],
                    [f"event:{index}", "model", event.model],
                    [f"event:{index}", "prompt_version", event.prompt_version],
                    [f"event:{index}", "started_at", event.started_at.isoformat()],
                    [f"event:{index}", "duration_ms", event.duration_ms],
                    [f"event:{index}", "input_pages", ",".join(map(str, event.input_pages))],
                    [f"event:{index}", "error", event.error],
                ]
            )
            rows.extend(
                [f"event:{index}", f"usage.{key}", value] for key, value in event.usage.items()
            )
        return rows

    @staticmethod
    def _confirmed_pages(run: Run) -> str | None:
        if run.location is None or not run.location.confirmed:
            return None
        return f"{run.location.page_start}-{run.location.page_end}"

    @staticmethod
    def _format_sheet(sheet: object) -> None:
        sheet.freeze_panes = "A2"  # type: ignore[attr-defined]
        sheet.auto_filter.ref = sheet.dimensions  # type: ignore[attr-defined]
        for cell in sheet[1]:  # type: ignore[index]
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for column in sheet.columns:  # type: ignore[attr-defined]
            width = min(
                60,
                max(12, max(len(str(cell.value or "")) for cell in column) + 2),
            )
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width  # type: ignore[attr-defined]


__all__ = ["ExcelExportError", "ExcelExporter"]
